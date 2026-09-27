from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "output" / "weekly_report" / "202237923_백승민_주간보고_2026-09-19.docx"
CONCEPT_IMAGE = ROOT / "tmp" / "pdfs" / "plan_render" / "page-11.png"

FONT_NAME = "Malgun Gothic"
LABEL_FILL = "DCE6F4"
LIGHT_BORDER = "B7C3D0"
BLACK = "000000"


def set_run_font(run, size: float, *, bold: bool = False, color: str = BLACK) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), FONT_NAME)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_margins(cell, *, top: int = 80, start: int = 110, bottom: int = 80, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, *, outer_size: int = 10, inside_size: int = 4, vertical: bool = True) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    border_specs = {
        "top": (BLACK, outer_size, "single"),
        "left": (BLACK, outer_size, "nil" if not vertical else "single"),
        "bottom": (BLACK, outer_size, "single"),
        "right": (BLACK, outer_size, "nil" if not vertical else "single"),
        "insideH": (LIGHT_BORDER, inside_size, "single"),
        "insideV": (LIGHT_BORDER, inside_size, "nil" if not vertical else "single"),
    }
    for edge, (color, size, style) in border_specs.items():
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), style)
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_spacing(paragraph, *, before: float = 0, after: float = 0, line: float = 1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def clear_cell(cell) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_spacing(paragraph)


def write_cell(
    cell,
    text: str,
    *,
    size: float = 8.7,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    v_align=WD_ALIGN_VERTICAL.CENTER,
) -> None:
    clear_cell(cell)
    cell.vertical_alignment = v_align
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold)


def add_body_line(cell, text: str, *, bold_prefix: str | None = None, after: float = 1.5) -> None:
    paragraph = cell.add_paragraph()
    set_paragraph_spacing(paragraph, after=after, line=1.0)
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        set_run_font(prefix_run, 8.5, bold=True)
        body_run = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(body_run, 8.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 8.5)


def set_row_height(row, height_cm: float) -> None:
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def add_spacer(document, height_pt: float) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, after=0)
    paragraph.paragraph_format.line_spacing = Pt(height_pt)
    run = paragraph.add_run(" ")
    set_run_font(run, 1)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(8.7)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    doc.core_properties.title = "프로젝트 주간 보고"
    doc.core_properties.subject = "1인칭 버블보블 프로젝트 3주차 진행 보고"
    doc.core_properties.author = "백승민"

    header = doc.add_table(rows=2, cols=6)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    widths = [5.7, 1.8, 1.6, 1.0, 2.0, 2.9]
    for row in header.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell, top=35, bottom=35, start=70, end=70)
    set_row_height(header.rows[0], 0.68)
    set_row_height(header.rows[1], 0.68)
    title_cell = header.cell(0, 0).merge(header.cell(1, 0))
    write_cell(title_cell, "프로젝트 주간 보고", size=12.2, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(0, 1), "날짜", size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(0, 2), "09", size=8.7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(0, 3), "월", size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(0, 4), "19", size=8.7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(0, 5), "일", size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(1, 1), "이름", size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    name_cell = header.cell(1, 2).merge(header.cell(1, 3))
    write_cell(name_cell, "백승민", size=8.7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(1, 4), "학번", size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(header.cell(1, 5), "202237923", size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in (header.cell(0, 1), header.cell(1, 1), header.cell(1, 4)):
        set_cell_shading(cell, LABEL_FILL)
    set_table_borders(header, outer_size=10, inside_size=4, vertical=False)

    add_spacer(doc, 4)

    project = doc.add_table(rows=1, cols=2)
    project.alignment = WD_TABLE_ALIGNMENT.CENTER
    project.autofit = False
    set_cell_width(project.cell(0, 0), 3.5)
    set_cell_width(project.cell(0, 1), 11.5)
    set_cell_shading(project.cell(0, 0), LABEL_FILL)
    set_cell_margins(project.cell(0, 0), top=70, bottom=70, start=80, end=80)
    set_cell_margins(project.cell(0, 1), top=70, bottom=70, start=130, end=130)
    write_cell(project.cell(0, 0), "프로젝트명", size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(project.cell(0, 1), "1인칭 버블보블 (Bubble Bobble FPS)", size=8.7, bold=True)
    set_row_height(project.rows[0], 0.9)
    set_table_borders(project, outer_size=8, inside_size=4, vertical=True)

    add_spacer(doc, 4)

    body = doc.add_table(rows=3, cols=2)
    body.alignment = WD_TABLE_ALIGNMENT.CENTER
    body.autofit = False
    for row in body.rows:
        set_cell_width(row.cells[0], 3.5)
        set_cell_width(row.cells[1], 11.5)
        set_cell_shading(row.cells[0], LABEL_FILL)
        set_cell_margins(row.cells[0], top=90, bottom=90, start=80, end=80)
        set_cell_margins(row.cells[1], top=110, bottom=100, start=140, end=140)
    set_row_height(body.rows[0], 7.8)
    set_row_height(body.rows[1], 6.2)
    set_row_height(body.rows[2], 4.5)
    write_cell(body.cell(0, 0), "작업 진행 내용", size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(body.cell(1, 0), "이미지/자료\n첨부", size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(body.cell(2, 0), "느낀 점", size=8.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_table_borders(body, outer_size=8, inside_size=5, vertical=True)

    progress = body.cell(0, 1)
    clear_cell(progress)
    intro = progress.paragraphs[0]
    intro_run = intro.add_run("기획서 3주차 목표인 포획 인터페이스와 Enemy 상태 구조를 중심으로 구현했다.")
    set_run_font(intro_run, 8.6, bold=True)
    set_paragraph_spacing(intro, after=3, line=1.0)
    progress_lines = [
        "- Bubble 상태를 Flying, Floating, Captured, Bursting으로 분리하고 비행 후 상승과 수명 종료 처리를 추가했다.",
        "- ICapturable 구현을 Capturable 컴포넌트로 분리했다. 포획 시 Rigidbody를 고정하고 적이 버블 위치를 따라가도록 구성했다.",
        "- Enemy 상태를 Idle, Chase, Attack, Captured, Dead로 정리하고 거리 기반 추적과 공격 전환을 연결했다.",
        "- PlayerHealth와 HP UI를 추가해 몬스터 공격이 체력에 반영되는 흐름을 확인했다.",
        "- 로비, 로딩, 전투 장면과 GameManager 상태 전환을 연결해 기본 게임 루프를 구성 중이다.",
        "- 포획 상태 전환은 동작하지만 Escape와 총기 Pop의 결과를 분리하는 작업은 남아 있다.",
    ]
    for line in progress_lines:
        add_body_line(progress, line, after=1.4)

    attachment = body.cell(1, 1)
    clear_cell(attachment)
    evidence_title = attachment.paragraphs[0]
    evidence_title_run = evidence_title.add_run("기획 비주얼과 구현 근거")
    set_run_font(evidence_title_run, 8.6, bold=True)
    set_paragraph_spacing(evidence_title, after=2)

    nested = attachment.add_table(rows=1, cols=2)
    nested.alignment = WD_TABLE_ALIGNMENT.CENTER
    nested.autofit = False
    set_cell_width(nested.cell(0, 0), 5.25)
    set_cell_width(nested.cell(0, 1), 5.45)
    set_cell_margins(nested.cell(0, 0), top=30, bottom=30, start=30, end=80)
    set_cell_margins(nested.cell(0, 1), top=30, bottom=30, start=80, end=30)
    set_table_borders(nested, outer_size=0, inside_size=0, vertical=False)

    image_cell = nested.cell(0, 0)
    clear_cell(image_cell)
    image_paragraph = image_cell.paragraphs[0]
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_run = image_paragraph.add_run()
    image_run.add_picture(str(CONCEPT_IMAGE), width=Cm(5.0))
    set_paragraph_spacing(image_paragraph, after=1)
    caption = image_cell.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption.add_run("기획서 비주얼 콘셉트")
    set_run_font(caption_run, 7.2, color="505050")
    set_paragraph_spacing(caption, after=0)

    evidence_cell = nested.cell(0, 1)
    clear_cell(evidence_cell)
    evidence_lines = [
        "상태 흐름",
        "Bubble  Flying > Floating / Captured > Bursting",
        "Enemy  Idle > Chase > Attack > Captured > Dead",
        "주요 파일",
        "Bubble.cs, Capturable.cs, BasicMonster.cs, MonsterBase.cs, PlayerHealth.cs, PlayerHealthUI.cs",
        "Git 기록  6f8fefd, 261909c, 47f07b9, 138b14a, 19f8a7c",
    ]
    for index, line in enumerate(evidence_lines):
        paragraph = evidence_cell.paragraphs[0] if index == 0 else evidence_cell.add_paragraph()
        set_paragraph_spacing(paragraph, after=1.2, line=1.0)
        run = paragraph.add_run(line)
        set_run_font(run, 7.6 if index not in (0, 3) else 8.0, bold=index in (0, 3))

    reflection = body.cell(2, 1)
    clear_cell(reflection)
    reflection_text = (
        "포획 기능을 몬스터 안에 직접 넣지 않고 Capturable 컴포넌트로 분리하니 버블과 적 로직의 역할이 명확해졌다. "
        "다만 현재는 버블 수명이 끝나도 적 사망으로 이어져 기획서의 Escape 규칙과 총기 Pop이 구분되지 않는다. "
        "다음 주에는 수명 종료 시 탈출, 총기 Pop 시 처치가 되도록 상태 전이를 나누고 OverlapSphere 기반 Chain을 구현할 계획이다."
    )
    reflection_run = reflection.paragraphs[0].add_run(reflection_text)
    set_run_font(reflection_run, 8.7)
    set_paragraph_spacing(reflection.paragraphs[0], after=0, line=1.15)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
